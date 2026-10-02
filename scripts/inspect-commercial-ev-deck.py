#!/usr/bin/env python3
import base64
from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
src = ROOT / "src/generated/commercial-ev-assets/Manash-Protim-Deori-Commercial-EV-Growth-Advisory-Case.pptx.01.b64"
work = ROOT / ".tmp-commercial-ev-inspect"
work.mkdir(exist_ok=True)
pptx = work / "deck.pptx"
pptx.write_bytes(base64.b64decode(src.read_text(encoding="utf-8").strip()))
prs = Presentation(pptx)

markers = [
    "NITI Aayog/WRI", "Commercial EV registration share", "Adoption can outrun public charging",
    "Illustrative 5-year TCO per km", "Verified vehicle inputs", "Recurring revenue potential",
    "Priority 1", "0-30 DAYS", "Six challenged figures", "claim reconciliation",
]

for i, slide in enumerate(prs.slides, start=1):
    texts = []
    for j, s in enumerate(slide.shapes):
        if getattr(s, "has_text_frame", False):
            t = " ".join((s.text or "").replace("\n"," ").split())
            if t:
                texts.append((j, s, t))
    slide_text = " || ".join(t for _,_,t in texts)
    if any(m.lower() in slide_text.lower() for m in markers):
        print(f"\n===== SLIDE {i} =====")
        for j, s, t in texts:
            print(f"[{j}] x={s.left/914400:.3f} y={s.top/914400:.3f} w={s.width/914400:.3f} h={s.height/914400:.3f} :: {t}")
